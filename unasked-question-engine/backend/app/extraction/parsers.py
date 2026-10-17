"""File parsers. Each returns an ``ExtractedDocument`` with source-referenced chunks.

All parsing happens on the backend. Failures raise ``ExtractionError`` with a
user-safe message; nothing here ever executes document content.
"""

from __future__ import annotations

import csv
import io
import logging
import re
import statistics
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from .chunking import MAX_CHUNK_CHARS, Block, pack_blocks, paragraph_blocks
from .models import Chunk, ExtractedDocument, ExtractionError, SourceRef

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS: dict[str, str] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "txt",
    ".md": "markdown",
    ".markdown": "markdown",
    ".csv": "csv",
}
MAX_DOCX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MIN_PAGE_TEXT_CHARS = 20


@dataclass
class Section:
    heading: str | None
    page: int | None
    blocks: list[Block]
    kind: str = "text"


@dataclass
class _Assembler:
    doc_id: str
    filename: str
    doc_type: str
    chunks: list[Chunk] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def add_chunk(self, text: str, source: SourceRef, kind: str = "text") -> None:
        self.chunks.append(
            Chunk(
                chunk_id=f"{self.doc_id}-c{len(self.chunks) + 1}",
                doc_id=self.doc_id,
                text=text,
                source=source,
                kind=kind,  # type: ignore[arg-type]
            )
        )

    def add_sections(self, sections: list[Section]) -> None:
        for section in sections:
            for block in pack_blocks(section.blocks):
                self.add_chunk(
                    block.text,
                    SourceRef(
                        filename=self.filename,
                        page=section.page,
                        heading=section.heading,
                        line_start=block.line_start,
                        line_end=block.line_end,
                    ),
                    section.kind,
                )

    def finish(self, metadata: dict | None = None) -> ExtractedDocument:
        if not self.chunks:
            raise ExtractionError(
                "empty_document", f"'{self.filename}' contains no readable text."
            )
        return ExtractedDocument(
            doc_id=self.doc_id,
            filename=self.filename,
            doc_type=self.doc_type,
            char_count=sum(len(c.text) for c in self.chunks),
            chunks=self.chunks,
            warnings=self.warnings,
            metadata=metadata or {},
        )


# --------------------------------------------------------------------------- text
def decode_bytes(data: bytes, filename: str) -> tuple[str, list[str]]:
    if b"\x00" in data[:8192]:
        raise ExtractionError(
            "binary_content", f"'{filename}' looks like a binary file, not text."
        )
    try:
        return data.decode("utf-8-sig"), []
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace"), [
            "File is not valid UTF-8; decoded as Windows-1252, so some characters may be wrong."
        ]


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_FENCE_RE = re.compile(r"^\s*(```|~~~)")


def _markdown_sections(lines: list[str]) -> list[Section]:
    sections: list[Section] = []
    stack: list[tuple[int, str]] = []
    current: list[str] = []
    start_line = 1
    in_fence = False

    def flush() -> None:
        blocks = paragraph_blocks(current, start_line)
        if blocks:
            heading = " > ".join(title for _, title in stack) or None
            sections.append(Section(heading, None, blocks))

    for index, line in enumerate(lines, start=1):
        if _FENCE_RE.match(line):
            in_fence = not in_fence
        match = None if in_fence else _HEADING_RE.match(line)
        if match:
            flush()
            level, title = len(match.group(1)), match.group(2).strip()
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, title))
            current, start_line = [], index + 1
        else:
            current.append(line)
    flush()
    return sections


def parse_text(text: str, filename: str, doc_id: str, *, markdown: bool = False,
               kind: str = "text") -> ExtractedDocument:
    if not text.strip():
        raise ExtractionError("empty_document", f"'{filename}' is empty.")
    lines = text.splitlines()
    asm = _Assembler(doc_id, filename, "markdown" if markdown else "txt")
    if markdown:
        sections = _markdown_sections(lines)
    else:
        sections = [Section(None, None, paragraph_blocks(lines, 1))]
    for section in sections:
        section.kind = kind
    asm.add_sections(sections)
    return asm.finish({"lines": len(lines)})


# --------------------------------------------------------------------------- csv
def _to_float(value: str) -> float | None:
    try:
        return float(value.replace(",", "").strip())
    except ValueError:
        return None


def _profile_text(filename: str, header: list[str], data: list[tuple[int, list[str]]],
                  ragged: int) -> str:
    lines = [
        f"Dataset profile for {filename}",
        f"Columns ({len(header)}): " + ", ".join(h or "(blank)" for h in header),
        f"Data rows: {len(data)}",
    ]
    if ragged:
        lines.append(f"Rows whose cell count differs from the header: {ragged}")
    for col_index, name in enumerate(header):
        values = [row[col_index].strip() if col_index < len(row) else "" for _, row in data]
        non_empty = [v for v in values if v]
        parts = [f"{len(values) - len(non_empty)} empty of {len(values)}"]
        numbers = [n for n in (_to_float(v) for v in non_empty) if n is not None]
        if non_empty and len(numbers) / len(non_empty) >= 0.8:
            parts.append(
                f"numeric: min={min(numbers):g}, max={max(numbers):g}, "
                f"mean={statistics.fmean(numbers):.4g}"
            )
        else:
            parts.append(f"{len(set(non_empty))} distinct values")
        lines.append(f"Column '{name or '(blank)'}': " + "; ".join(parts))
    return "\n".join(lines)


def parse_csv(data: bytes, filename: str, doc_id: str, *, max_rows: int) -> ExtractedDocument:
    text, warnings = decode_bytes(data, filename)
    if not text.strip():
        raise ExtractionError("empty_document", f"'{filename}' is empty.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        rows = list(csv.reader(io.StringIO(text, newline=""), dialect))
    except csv.Error as exc:
        raise ExtractionError("malformed_csv", f"'{filename}' is not valid CSV: {exc}") from exc

    numbered = [(i, r) for i, r in enumerate(rows, start=1) if any(c.strip() for c in r)]
    if not numbered:
        raise ExtractionError("empty_document", f"'{filename}' contains no rows.")
    header_row, header = numbered[0]
    header = [h.strip() for h in header]
    data = numbered[1:]
    if not data:
        raise ExtractionError(
            "empty_document", f"'{filename}' has a header row but no data rows."
        )
    if len(data) > max_rows:
        warnings.append(f"Only the first {max_rows} of {len(data)} data rows were analysed.")
        data = data[:max_rows]
    ragged = sum(1 for _, r in data if len(r) != len(header))
    if ragged:
        warnings.append(f"{ragged} row(s) have a different number of cells than the header.")
    dup_headers = {h for h in header if h and header.count(h) > 1}
    if dup_headers:
        warnings.append("Duplicate column names: " + ", ".join(sorted(dup_headers)) + ".")

    asm = _Assembler(doc_id, filename, "csv", warnings=warnings)
    asm.add_chunk(
        _profile_text(filename, header, data, ragged),
        SourceRef(filename=filename, heading="Dataset profile",
                  row_start=header_row, row_end=data[-1][0]),
        "dataset_profile",
    )
    header_line = "Columns: " + " | ".join(header)
    header_line = header_line[:300]
    buffer: list[str] = []
    first_row = last_row = 0
    size = len(header_line)

    def flush() -> None:
        if buffer:
            asm.add_chunk(
                header_line + "\n" + "\n".join(buffer),
                SourceRef(filename=filename, row_start=first_row, row_end=last_row),
                "dataset_rows",
            )

    for row_no, row in data:
        cells = [
            f"{header[i] if i < len(header) and header[i] else f'col{i + 1}'}={cell.strip()[:200]}"
            for i, cell in enumerate(row)
            if cell.strip()
        ]
        line = f"Row {row_no}: " + "; ".join(cells) if cells else f"Row {row_no}: (all cells empty)"
        if buffer and size + len(line) + 1 > MAX_CHUNK_CHARS:
            flush()
            buffer, size = [], len(header_line)
        if not buffer:
            first_row = row_no
        buffer.append(line)
        last_row = row_no
        size += len(line) + 1
    flush()
    return asm.finish(
        {"columns": header, "data_rows": len(data),
         "row_numbering": "Spreadsheet-style: the header is the first non-empty row."}
    )


# --------------------------------------------------------------------------- pdf
def parse_pdf(path: Path, filename: str, doc_id: str, *, max_pages: int) -> ExtractedDocument:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractionError(
                "encrypted_pdf", f"'{filename}' is password-protected and cannot be read."
            )
        page_count = len(reader.pages)
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001 - pypdf raises many exception types
        logger.info("PDF open failed for %s: %s", type(exc).__name__, exc)
        raise ExtractionError(
            "malformed_pdf", f"'{filename}' could not be read as a PDF (file may be corrupt)."
        ) from exc

    asm = _Assembler(doc_id, filename, "pdf")
    if page_count > max_pages:
        asm.warnings.append(f"Only the first {max_pages} of {page_count} pages were analysed.")
    low_text_pages: list[int] = []
    total_chars = 0
    for number in range(1, min(page_count, max_pages) + 1):
        try:
            text = reader.pages[number - 1].extract_text() or ""
        except Exception as exc:  # noqa: BLE001
            asm.warnings.append(f"Page {number} could not be read ({type(exc).__name__}).")
            continue
        text = re.sub(r"[ \t]+\n", "\n", text)
        total_chars += len(text.strip())
        if len(text.strip()) < MIN_PAGE_TEXT_CHARS:
            low_text_pages.append(number)
        blocks = [Block(b.text) for b in paragraph_blocks(text.splitlines())]
        asm.add_sections([Section(None, number, blocks)])
    if total_chars == 0:
        raise ExtractionError(
            "no_extractable_text",
            f"'{filename}' has no extractable text (it may be a scanned image; OCR is not supported).",
        )
    if low_text_pages:
        shown = ", ".join(str(p) for p in low_text_pages[:10])
        asm.warnings.append(
            f"Little or no text on page(s) {shown}; they may be images or scans."
        )
    return asm.finish({"pages": page_count})


# --------------------------------------------------------------------------- docx
def parse_docx(path: Path, filename: str, doc_id: str) -> ExtractedDocument:
    try:
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ExtractionError(
                    "file_too_large", f"'{filename}' expands to an unreasonable size."
                )
    except zipfile.BadZipFile as exc:
        raise ExtractionError(
            "malformed_docx", f"'{filename}' is not a valid .docx file."
        ) from exc

    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(str(path))
        body_items = []
        for child in document.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                body_items.append(Paragraph(child, document))
            elif tag == "tbl":
                body_items.append(Table(child, document))
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001
        logger.info("DOCX open failed for %s: %s", type(exc).__name__, exc)
        raise ExtractionError(
            "malformed_docx", f"'{filename}' could not be read as a Word document."
        ) from exc

    asm = _Assembler(doc_id, filename, "docx")
    sections: list[Section] = []
    stack: list[tuple[int, str]] = []
    lines: list[str] = []

    def flush() -> None:
        blocks = paragraph_blocks(lines)
        if blocks:
            heading = " > ".join(t for _, t in stack) or None
            sections.append(Section(heading, None, [Block(b.text) for b in blocks]))

    for item in body_items:
        if isinstance(item, Paragraph):
            text = item.text.strip()
            style = (item.style.name if item.style is not None else "") or ""
            level_match = re.match(r"Heading\s*(\d+)$", style)
            if text and (level_match or style == "Title"):
                flush()
                lines = []
                level = int(level_match.group(1)) if level_match else 1
                while stack and stack[-1][0] >= level:
                    stack.pop()
                stack.append((level, text))
            else:
                lines.append(text)
                lines.append("")
        else:
            lines.append("")
            for row in item.rows:
                cells: list[str] = []
                for cell in row.cells:
                    cell_text = cell.text.strip().replace("\n", " ")
                    if not cells or cells[-1] != cell_text:
                        cells.append(cell_text)
                if any(cells):
                    lines.append(" | ".join(cells))
            lines.append("")
    flush()
    asm.add_sections(sections)
    return asm.finish({"tables": len(document.tables)})


# --------------------------------------------------------------------------- dispatch
def parse_file(path: Path, filename: str, doc_id: str, *, max_csv_rows: int = 20000,
               max_pdf_pages: int = 300) -> ExtractedDocument:
    extension = Path(filename).suffix.lower()
    doc_type = SUPPORTED_EXTENSIONS.get(extension)
    if doc_type is None:
        raise ExtractionError(
            "unsupported_type",
            f"'{filename}' has an unsupported type. Supported: "
            + ", ".join(sorted(SUPPORTED_EXTENSIONS)) + ".",
        )
    head = path.read_bytes()[:1024] if doc_type in {"pdf", "docx"} else b""
    if doc_type == "pdf" and b"%PDF" not in head:
        raise ExtractionError("content_mismatch", f"'{filename}' is not a PDF despite its extension.")
    if doc_type == "docx" and not head.startswith(b"PK"):
        raise ExtractionError("content_mismatch", f"'{filename}' is not a .docx despite its extension.")

    if doc_type == "pdf":
        return parse_pdf(path, filename, doc_id, max_pages=max_pdf_pages)
    if doc_type == "docx":
        return parse_docx(path, filename, doc_id)
    data = path.read_bytes()
    if doc_type == "csv":
        return parse_csv(data, filename, doc_id, max_rows=max_csv_rows)
    text, warnings = decode_bytes(data, filename)
    document = parse_text(text, filename, doc_id, markdown=doc_type == "markdown")
    document.warnings.extend(warnings)
    return document
